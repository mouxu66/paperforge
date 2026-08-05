"""学术诚信报告组装与渲染（学生「AI 使用声明 + 真实性报告」一页纸）。

设计原则（详见 docs/integrity-report-plan 讨论）：
- **证据性文档必须确定性生成**：报告里的每个数字都逐字来自数据库记录
  （``DepthReviewV4.reflection_result["analysis_v2"]``），禁止 LLM 转述——
  转述即引入幻觉风险。
- **无过程 provenance 就明说**：感悟报告场景学生交独立 docx，系统没有
  写作过程记录 → AI 使用声明走 ``post_hoc`` 模式（advisory 疑似度 + 免责），
  绝不让教师误以为系统能"检测"AI（与 ADR-009 的 advisory 红线一致）。
- **可信度标记进报告**：``llm_failed`` / ``effective_evidence_count`` 等
  诊断字段直接上报告，把「系统故障」与「学生报告质量」分开
  （对应 depth_eval_reflection 2026-08-05 埋点）。

对外 API：
    build_integrity_report(db, paper_id) -> dict | None   # 组装（无 IO 副作用）
    render_integrity_docx(report) -> bytes                # python-docx 渲染
    render_integrity_html(report) -> str                  # 打印版 HTML（浏览器存 PDF）

任何字段缺失都不抛异常（诊断是辅助信息，宁可缺失不能拖垮主链路）。
"""

from __future__ import annotations

import html
import logging
import re
import time
import uuid
import zipfile
from collections.abc import Callable
from datetime import datetime
from pathlib import Path

from sqlalchemy.orm import Session

from .models import DepthReviewV4
from .models import Paper as PaperORM
from .routers.common import as_dict

logger = logging.getLogger(__name__)

# ── 中文展示映射（仅报告展示用，不参与业务判定）──────────────────────
VERDICT_LABELS = {
    "well_done": "优秀",
    "needs_evidence": "证据不足",
    "needs_depth": "深度不足",
    "rewrite_required": "需重写",
}
AI_TIER_LABELS = {
    "human": "疑似人工写作",
    "uncertain": "无法判断",
    "likely_ai": "疑似 AI 生成",
}
SECTION_LABELS = {
    "q": "问题与目的",
    "tech": "技术方法",
    "exp": "实验部分",
    "reflection": "感想反思",
}

# 默认评分权重（与 reflection_pipeline.W 同步；analysis_v2 未携带 weights 时兜底）
DEFAULT_WEIGHTS = {
    "understanding_accuracy": 0.05,
    "analysis_depth": 0.15,
    "innovative_insights": 0.35,
    "evidence_support": 0.05,
    "fidelity": 0.05,
    "coverage": 0.35,
}

# 与 depth_eval_reflection.MIN_EVIDENCE_FOR_VALID_REVIEW 保持一致
MIN_EVIDENCE_FOR_VALID_REVIEW = 2


def _safe_bool(value, default: bool = False) -> bool:
    """严格布尔化：只有真实的数值/布尔才转，字符串等一律回默认值。

    与 reflection_pipeline._safe_bool 同语义——防止老库把 llm_failed 存成
    字符串 "false" 时 bool("false") == True 的陷阱。
    """
    if isinstance(value, (bool, int, float)):
        return bool(value)
    return default


def _default_weights() -> dict:
    """取评分权重：优先读 reflection_pipeline.W（单一事实源），失败用本地兜底。

    延迟导入：reflection_pipeline 会连带加载 fastembed 等重依赖，
    不能在模块顶层 import（会拖慢 app 启动并触发模型下载尝试）。
    """
    try:
        from .reflection_pipeline import W as pipeline_W

        if isinstance(pipeline_W, dict) and pipeline_W:
            return dict(pipeline_W)
    except Exception:  # noqa: BLE001 - 权重兜底 - 失败用本地常量
        pass
    return dict(DEFAULT_WEIGHTS)


def _fmt(v, nd: int = 2) -> str:
    """0~1 分数格式化；None/异常 → '—'。"""
    if v is None:
        return "—"
    try:
        return f"{float(v):.{nd}f}"
    except (TypeError, ValueError):
        return "—"


def _fmt_pct(v) -> str:
    """0~1 比率格式化；None/异常 → '—'。"""
    if v is None:
        return "—"
    try:
        return f"{float(v) * 100:.1f}%"
    except (TypeError, ValueError):
        return "—"


def build_integrity_report(db: Session, paper_id: str) -> dict | None:
    """从数据库组装一篇感悟报告的「AI 使用声明 + 真实性报告」。

    数据源：``DepthReviewV4(kind='report').reflection_result["analysis_v2"]``
    （reflection_pipeline.analyze_reflection_file 的完整结果）；老记录无
    ``analysis_v2`` 时回退到 ``reflection_result`` 顶层字段（4 维 scores 等）。

    返回 None 表示没有该报告的评审记录。
    """
    rev = (
        db.query(DepthReviewV4)
        .filter(DepthReviewV4.paper_id == paper_id, DepthReviewV4.kind == "report")
        .order_by(DepthReviewV4.created_at.desc())
        .first()
    )
    if rev is None:
        return None
    return _build_from_review(db, rev)


def _build_from_review(db: Session, rev: DepthReviewV4) -> dict:
    """由一条评审记录组装报告（单篇/批量导出共用，避免逐条重复查询）。"""
    rr = as_dict(rev.reflection_result)
    ext = as_dict(rr.get("analysis_v2"))  # 完整分析（6 维）；老记录可能为空 dict

    def pick(name, default=None):
        """优先取 analysis_v2，缺失时回退 reflection_result 顶层字段。"""
        if ext.get(name) is not None:
            return ext[name]
        if rr.get(name) is not None:
            return rr[name]
        return default

    # 报告 paper 行（标题兜底 / 绑定关系）
    paper = db.query(PaperORM).filter(PaperORM.id == rev.paper_id).first()
    student_id = pick("student_id") or ""
    student_name = pick("student_name") or ""

    # 原论文（只读引用，用于报告头部）
    bound = pick("bound_paper_id") or (paper.source_paper_id if paper else None)
    source_paper = None
    if bound:
        src_row = db.query(PaperORM).filter(PaperORM.id == bound).first()
        if src_row:
            source_paper = {
                "paper_id": src_row.id,
                "title": src_row.title or "",
                "author": ", ".join(src_row.authors or []),
                "year": src_row.year or 0,
                "source": src_row.source or "",
            }

    # ① AI 使用声明（post_hoc：无过程 provenance，advisory 弱信号）
    advisory = None
    ai_like = pick("ai_likelihood")
    if ai_like is not None:
        ai_tier = pick("ai_likelihood_tier") or ""
        # 注意：analysis_v2 未持久化 confidence（ai_likelihood 模块恒为 "low"），
        # 不在这里编造字段——advisory 只列真实存在的数据。
        advisory = {
            "ai_likelihood": ai_like,
            "tier": ai_tier,
            "tier_label": AI_TIER_LABELS.get(ai_tier, ai_tier),
            "signals": pick("ai_likelihood_signals") or {},
            "note": pick("ai_likelihood_note") or "",
        }
    ai_use = {
        "mode": "post_hoc",
        "provenance_available": False,
        "declaration_text": (
            "本感悟报告为独立提交的文档，系统未记录其写作过程，"
            "因此无法提供基于过程溯源的 AI 使用声明。"
        ),
        "advisory": advisory,
        "advisory_note": (
            "AI 生成疑似度为文本特征的事后弱信号，仅供参考、非证据，"
            "不可作为处分依据；如有疑问请结合写作过程与面谈核实。"
        ),
    }

    # ② 真实性核验（报告↔原论文）
    authenticity = {
        "fidelity": pick("fidelity"),
        "fidelity_status": pick("fidelity_status") or "",
        "coverage": pick("coverage"),
        "coverage_status": pick("coverage_status") or "",
        "coverage_covered": pick("coverage_covered") or [],
        "coverage_uncovered": pick("coverage_uncovered") or [],
        "copy_ratio": pick("copy_ratio"),
        "copy_sentences": pick("copy_sentences") or [],
        "stray_claims": pick("stray_claims") or [],
        "anchors": pick("fidelity_anchors") or [],
        "sections_present": [SECTION_LABELS.get(s, s) for s in (pick("sections_present") or [])],
    }

    # ③ 评分 + 可信度诊断（2026-08-05 埋点：LLM 故障与报告质量分开）
    raw_scores = ext.get("scores") or rr.get("scores") or {}
    if not isinstance(raw_scores, dict):
        raw_scores = {}
    scores = {
        k: round(float(v), 4) if isinstance(v, (int, float)) else v
        for k, v in raw_scores.items()
        if v is not None
    }
    weights = ext.get("weights") if isinstance(ext.get("weights"), dict) else {}
    # 旧版本在 LLM 故障时可能只写 status='failed' 而没有 reflection_result。
    # 这类记录也必须视为不可信，不能因为缺诊断字段而回退到历史分数。
    record_failed = rev.status in {"failed", "timed_out"}
    llm_failed = _safe_bool(pick("llm_failed", False)) or record_failed
    parse_failed = _safe_bool(pick("parse_failed", False))
    truncated = _safe_bool(pick("truncated", False))
    ev_count = pick("effective_evidence_count")

    trust_warnings: list[str] = []
    if llm_failed:
        trust_warnings.append(
            "本次评审中 LLM 调用失败/空返回，评分不可信（系统故障，不代表报告质量）"
        )
    if parse_failed:
        trust_warnings.append("报告解析失败，评审结果不完整")
    if ev_count is not None and ev_count < MIN_EVIDENCE_FOR_VALID_REVIEW:
        trust_warnings.append(
            f"有效证据不足 {MIN_EVIDENCE_FOR_VALID_REVIEW} 条，四维分数已按硬规则封顶（R1）"
        )
    if truncated:
        trust_warnings.append("报告超出模型输入上限，评审仅基于报告头尾部分")

    verdict = "llm_failed" if llm_failed else (pick("verdict") or "needs_evidence")
    if llm_failed:
        # 故障记录不能向报告/前端泄露任何历史评分或伪造结论，
        # 包括 fidelity/coverage 这类交叉评分。
        scores = {}
        authenticity["fidelity"] = None
        authenticity["coverage"] = None
        authenticity["copy_ratio"] = None
        authenticity["copy_sentences"] = []
        authenticity["stray_claims"] = []
    scoring = {
        "scores": scores,
        "weights": weights or _default_weights(),
        "average": None if llm_failed else pick("average"),
        "verdict": verdict,
        "verdict_label": VERDICT_LABELS.get(verdict, verdict),
        "verdict_reason": rr.get("verdict_reason") or "",
        "trusted": not llm_failed,
        "trust_warnings": trust_warnings,
        "hardcoded_overrides": pick("hardcoded_overrides") or [],
        "effective_evidence_count": ev_count,
        "evidence_rejections": pick("evidence_rejections") or {},
    }

    return {
        "report_type": "reflection",
        "paper_id": rev.paper_id,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "evaluated_at": (
            rev.completed_at.isoformat(timespec="seconds") if rev.completed_at else ""
        ),
        "student": {"id": student_id, "name": student_name},
        "title": pick("paper_title") or (paper.title if paper else rev.paper_id),
        "report_chars": pick("report_chars"),
        "paper_chars": pick("paper_chars"),
        "source_paper": source_paper,
        "ai_use": ai_use,
        "authenticity": authenticity,
        "scoring": scoring,
        "comments": None,  # LLM 教师评语草稿（可选，后续迭代）
    }


# ===========================================================================
# 渲染：docx（python-docx，主归档格式）
# ===========================================================================
def render_integrity_docx(report: dict) -> bytes:
    """把报告渲染为 .docx 字节流（中文宋体，一页纸版式）。"""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    NAVY = (0x1F, 0x4E, 0x79)
    RED = (0xC0, 0x39, 0x2B)

    def _set_cn(run, size: float = 10.5, bold: bool = False, color=None):
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor(*color)
        run.font.name = "Calibri"
        rpr = run._element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")

    def para(text, size: float = 10.5, bold: bool = False, color=None, space_before=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        _set_cn(r, size=size, bold=bold, color=color)
        return p

    def kv(label: str, value):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"{label}：")
        _set_cn(r1, size=10.5, bold=True)
        r2 = p.add_run("—" if value in (None, "") else str(value))
        _set_cn(r2, size=10.5)
        return p

    def section(title: str):
        return para(title, size=13, bold=True, color=NAVY, space_before=12)

    def bullet(text: str):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        _set_cn(r, size=10.5)
        return p

    # ── 头部 ──
    para("AI 使用声明与真实性报告", size=18, bold=True, color=NAVY)
    para("PaperForge 自动生成 · 数据可溯源至本地数据库 · 供教学参考", size=9)

    student = report.get("student") or {}
    kv("学号", student.get("id"))
    kv("姓名", student.get("name"))
    kv("报告", report.get("title"))
    kv("原论文", (report.get("source_paper") or {}).get("title") or "—")
    kv("评审时间", report.get("evaluated_at") or "—")
    kv("生成时间", report.get("generated_at") or "—")

    # ── ① AI 使用声明 ──
    section("一、AI 使用声明")
    ai_use = report.get("ai_use") or {}
    para(ai_use.get("declaration_text", ""), size=10.5)
    advisory = ai_use.get("advisory")
    if advisory:
        para(
            f"AI 生成疑似度：{_fmt(advisory.get('ai_likelihood'))}"
            f"（{advisory.get('tier_label') or advisory.get('tier') or '—'}）",
            size=10.5,
        )
        para(ai_use.get("advisory_note", ""), size=9, color=RED)

    # ── ② 真实性核验 ──
    section("二、真实性核验（报告 ↔ 原论文）")
    auth = report.get("authenticity") or {}
    kv("忠实度 fidelity", f"{_fmt(auth.get('fidelity'))}（{auth.get('fidelity_status') or '—'}）")
    kv("覆盖度 coverage", f"{_fmt(auth.get('coverage'))}（{auth.get('coverage_status') or '—'}）")
    kv("照抄率 copy_ratio", _fmt_pct(auth.get("copy_ratio")))
    if auth.get("coverage_uncovered"):
        para("未覆盖的核心要点：", size=10.5, bold=True)
        for item in list(auth["coverage_uncovered"])[:5]:
            bullet(str(item))
    if auth.get("stray_claims"):
        para(f"无出处论点 {len(auth['stray_claims'])} 条：", size=10.5, bold=True)
        for item in list(auth["stray_claims"])[:3]:
            bullet(str(item))

    # ── ③ 评分 ──
    section("三、评分（附权重，平均分 = Σ 分数 × 权重）")
    scoring = report.get("scoring") or {}
    scores = scoring.get("scores") or {}
    weights = scoring.get("weights") or {}
    if scores:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, h in enumerate(("维度", "分数", "权重")):
            r = hdr[i].paragraphs[0].add_run(h)
            _set_cn(r, size=10, bold=True)
        for key, val in scores.items():
            cells = table.add_row().cells
            cells[0].text = str(key)
            cells[1].text = _fmt(val)
            cells[2].text = _fmt(weights.get(key), 2)
            for cell in cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        _set_cn(r, size=10)
    kv("平均分", _fmt(scoring.get("average")))
    kv(
        "评审结论",
        f"{scoring.get('verdict_label') or scoring.get('verdict')}",
    )
    reason = scoring.get("verdict_reason")
    if reason:
        kv("结论依据", reason)
    for w in scoring.get("trust_warnings") or []:
        para(f"⚠ {w}", size=9, color=RED)
    for ov in scoring.get("hardcoded_overrides") or []:
        para(f"· {ov}", size=9)

    # ── 页脚 ──
    para("", size=6)
    para(
        "本报告由 PaperForge 根据本地评审记录自动生成，供教学参考；"
        "AI 生成疑似度为事后弱信号，不构成对学生的处分依据。",
        size=9,
    )

    from io import BytesIO

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ===========================================================================
# 渲染：打印版 HTML（浏览器 window.print() 存 PDF，零新依赖）
# ===========================================================================
def render_integrity_html(report: dict) -> str:
    """把报告渲染为打印友好的 HTML 片段（A4 版式，可直接保存为 PDF）。"""
    esc = html.escape

    def kv_row(label: str, value) -> str:
        return f"<tr><th>{esc(str(label))}</th><td>{esc('—' if value in (None, '') else str(value))}</td></tr>"

    student = report.get("student") or {}
    src = report.get("source_paper") or {}
    ai_use = report.get("ai_use") or {}
    auth = report.get("authenticity") or {}
    scoring = report.get("scoring") or {}
    scores = scoring.get("scores") or {}
    weights = scoring.get("weights") or {}

    parts: list[str] = []
    parts.append("<!DOCTYPE html><html lang='zh-CN'><head><meta charset='utf-8'>")
    parts.append(
        "<style>"
        "body{font-family:'Microsoft YaHei','SimSun',sans-serif;color:#222;"
        "max-width:720px;margin:24px auto;padding:0 16px;font-size:13px;line-height:1.6}"
        "h1{font-size:20px;color:#1f4e79;border-bottom:2px solid #1f4e79;"
        "padding-bottom:6px}h2{font-size:15px;color:#1f4e79;margin-top:18px}"
        "table{border-collapse:collapse;width:100%;margin:6px 0}"
        "th,td{border:1px solid #ccc;padding:4px 8px;text-align:left;font-size:12px}"
        "th{background:#eef3fa;width:140px}.warn{color:#c0392b;font-size:12px}"
        ".foot{margin-top:18px;border-top:1px solid #ddd;padding-top:6px;"
        "font-size:11px;color:#777}@media print{body{margin:0}}"
        "</style></head><body>"
    )

    parts.append("<h1>AI 使用声明与真实性报告</h1>")
    parts.append(
        "<p style='font-size:11px;color:#777'>PaperForge 自动生成 · 数据可溯源 · 供教学参考</p>"
    )
    parts.append(
        "<table>"
        + kv_row("学号", student.get("id"))
        + kv_row("姓名", student.get("name"))
        + kv_row("报告", report.get("title"))
        + kv_row("原论文", src.get("title") or "—")
        + kv_row("评审时间", report.get("evaluated_at") or "—")
        + kv_row("生成时间", report.get("generated_at") or "—")
        + "</table>"
    )

    # ① AI 使用声明
    parts.append("<h2>一、AI 使用声明</h2>")
    parts.append(f"<p>{esc(ai_use.get('declaration_text', ''))}</p>")
    advisory = ai_use.get("advisory")
    if advisory:
        parts.append(
            f"<p>AI 生成疑似度：<b>{esc(_fmt(advisory.get('ai_likelihood')))}</b>"
            f"（{esc(advisory.get('tier_label') or advisory.get('tier') or '—')}）</p>"
        )
        parts.append(f"<p class='warn'>{esc(ai_use.get('advisory_note', ''))}</p>")

    # ② 真实性核验
    parts.append("<h2>二、真实性核验（报告 ↔ 原论文）</h2>")
    parts.append(
        "<table>"
        + kv_row(
            "忠实度 fidelity",
            f"{_fmt(auth.get('fidelity'))}（{auth.get('fidelity_status') or '—'}）",
        )
        + kv_row(
            "覆盖度 coverage",
            f"{_fmt(auth.get('coverage'))}（{auth.get('coverage_status') or '—'}）",
        )
        + kv_row("照抄率 copy_ratio", _fmt_pct(auth.get("copy_ratio")))
        + "</table>"
    )
    if auth.get("coverage_uncovered"):
        parts.append(
            f"<p><b>未覆盖的核心要点：</b>{esc('；'.join(str(x) for x in auth['coverage_uncovered'][:5]))}</p>"
        )
    if auth.get("stray_claims"):
        parts.append(
            f"<p><b>无出处论点 {len(auth['stray_claims'])} 条：</b>{esc('；'.join(str(x) for x in auth['stray_claims'][:3]))}</p>"
        )

    # ③ 评分
    parts.append("<h2>三、评分（平均分 = Σ 分数 × 权重）</h2>")
    if scores:
        rows = "".join(
            f"<tr><td>{esc(str(k))}</td><td>{esc(_fmt(v))}</td>"
            f"<td>{esc(_fmt(weights.get(k), 2))}</td></tr>"
            for k, v in scores.items()
        )
        parts.append("<table><tr><th>维度</th><th>分数</th><th>权重</th></tr>" + rows + "</table>")
    parts.append(
        "<table>"
        + kv_row("平均分", _fmt(scoring.get("average")))
        + kv_row("评审结论", scoring.get("verdict_label") or scoring.get("verdict"))
        + ("</table>")
    )
    if scoring.get("verdict_reason"):
        parts.append(f"<p>结论依据：{esc(scoring['verdict_reason'])}</p>")
    for w in scoring.get("trust_warnings") or []:
        parts.append(f"<p class='warn'>⚠ {esc(w)}</p>")
    for ov in scoring.get("hardcoded_overrides") or []:
        parts.append(f"<p class='warn'>· {esc(str(ov))}</p>")

    parts.append(
        "<p class='foot'>本报告由 PaperForge 根据本地评审记录自动生成，供教学参考；"
        "AI 生成疑似度为事后弱信号，不构成对学生的处分依据。</p>"
    )
    parts.append("</body></html>")
    return "".join(parts)


# ===========================================================================
# 批量：全班诚信报告打包 zip（每生一份 docx）
# ===========================================================================
def render_integrity_batch_zip(
    db: Session,
    progress_cb: Callable[[int, int], None] | None = None,
) -> Path:
    """把全部感悟报告评审记录打包为 zip（每生一份 docx），返回 zip 文件路径。

    - 复用 render_integrity_docx；单篇失败不影响整批（记录 warning）。
    - 文件名：{学号}_{姓名}_诚信报告.docx（学号/姓名均做消毒，防 zip 注入）。
    - 无任何报告记录时生成一个含 README 说明的空 zip，避免教师拿到看不懂的空包。
    - zip 落在 DATA_DIR/exports/ 下；超过 24h 的历史 zip 自动清理。
    """
    from .database import DATA_DIR

    records = (
        db.query(DepthReviewV4)
        .filter(DepthReviewV4.kind == "report")
        .order_by(DepthReviewV4.created_at.desc())
        .all()
    )
    total = len(records)
    out_dir = DATA_DIR / "exports"
    out_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    zip_path = out_dir / f"integrity_reports_{ts}_{uuid.uuid4().hex[:8]}.zip"

    written = 0
    failed = 0
    seen_paper_ids: set[str] = (
        set()
    )  # 同一报告被重评多次时只保留最新一条（查询已按 created_at 倒序）
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        if not records:
            zf.writestr(
                "README.txt",
                "当前没有感悟报告评审记录，无法导出诚信报告。\n",
            )
        for i, rev in enumerate(records):
            if rev.paper_id in seen_paper_ids:
                # 重评产生的旧记录：跳过，否则 zip 内同名文件会被静默覆盖
                continue
            seen_paper_ids.add(rev.paper_id)
            try:
                report = _build_from_review(db, rev)
                data = render_integrity_docx(report)
                student = report.get("student") or {}
                sid = re.sub(
                    r"[^A-Za-z0-9_-]", "", student.get("id") or report.get("paper_id") or ""
                )
                sid = sid or "report"
                # 姓名可能含 /\:*?"<> 等非法字符，消毒后进 zip 文件名
                name = re.sub(r"[\\/:*?\"<>|\x00-\x1f]+", "_", student.get("name") or "").strip("_")
                fname = f"{sid}_{name}_诚信报告.docx" if name else f"{sid}_诚信报告.docx"
                zf.writestr(fname, data)
                written += 1
            except Exception:  # noqa: BLE001 - 批量打包 - 单篇失败不拖垮整批
                failed += 1
                logger.warning("批量诚信报告打包失败: paper=%s", rev.paper_id, exc_info=True)
            if progress_cb:
                progress_cb(i + 1, total)

    # 清理超过 24h 的历史 zip：下载端点不再删除文件（支持失败重试/重复下载），
    # 由这里统一回收，避免 exports 目录无限堆积。
    try:
        cutoff = time.time() - 86400
        for old in out_dir.glob("integrity_reports_*.zip"):
            if old.stat().st_mtime < cutoff:
                old.unlink(missing_ok=True)
    except OSError:
        pass

    logger.info(
        "全班诚信报告打包完成: %s（成功 %d，失败 %d）→ %s", zip_path.name, written, failed, zip_path
    )
    return zip_path
